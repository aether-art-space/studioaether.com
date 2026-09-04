#!/usr/bin/env python3
"""Capture a read-only SEO and URL baseline from the current production site.

The URL list comes from the committed migration inventory. The report stores
metadata and payload hashes, not full HTML snapshots, so it remains compact
while preserving enough evidence to compare the replacement site later.
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import tempfile
import time
import csv
from collections import Counter
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import urldefrag, urljoin, urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INVENTORY = ROOT / "docs/source/studioaether_wix_migration_inventory_2026-09-04.xlsx"
OUTPUT_DIR = ROOT / "docs/baseline"
SITE_ORIGIN = "https://www.studioaether.com"
USER_AGENT = "studioaether-v2-migration-baseline/1.0 (+read-only URL audit)"
TIMEOUT_SECONDS = 30
REQUEST_DELAY_SECONDS = 0.15


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def normalize_url(value: str | None, base: str = SITE_ORIGIN) -> str | None:
    if not value:
        return None
    value = value.strip()
    if value.startswith(("mailto:", "tel:", "javascript:", "#")):
        return None
    absolute = urljoin(base, value)
    absolute, _fragment = urldefrag(absolute)
    parsed = urlparse(absolute)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return None
    path = parsed.path or "/"
    if path != "/" and path.endswith("/"):
        path = path[:-1]
    return f"{parsed.scheme}://{parsed.netloc}{path}" + (f"?{parsed.query}" if parsed.query else "")


class PageParser(HTMLParser):
    def __init__(self, page_url: str) -> None:
        super().__init__(convert_charrefs=True)
        self.page_url = page_url
        self.title_parts: list[str] = []
        self.description: str | None = None
        self.robots: str | None = None
        self.googlebot: str | None = None
        self.canonical: str | None = None
        self.alternates: list[dict[str, str]] = []
        self.lang: str | None = None
        self.headings: dict[str, list[str]] = {"h1": [], "h2": [], "h3": []}
        self.links: list[dict[str, str]] = []
        self.images: list[dict[str, str | None]] = []
        self.json_ld_parts: list[list[str]] = []
        self._text_target: tuple[str, list[str]] | None = None
        self._current_heading: tuple[str, list[str]] | None = None
        self._current_json_ld: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_map = {key.lower(): value for key, value in attrs}
        tag = tag.lower()
        if tag == "html":
            self.lang = attrs_map.get("lang")
        elif tag == "title":
            self._text_target = ("title", self.title_parts)
        elif tag in self.headings:
            parts: list[str] = []
            self._current_heading = (tag, parts)
        elif tag == "meta":
            name = (attrs_map.get("name") or "").lower()
            if name == "description":
                self.description = attrs_map.get("content")
            elif name == "robots":
                self.robots = attrs_map.get("content")
            elif name == "googlebot":
                self.googlebot = attrs_map.get("content")
        elif tag == "link":
            rel = {item.strip().lower() for item in (attrs_map.get("rel") or "").split()}
            href = attrs_map.get("href")
            if "canonical" in rel and href:
                self.canonical = normalize_url(href, self.page_url)
            if "alternate" in rel and attrs_map.get("hreflang") and href:
                self.alternates.append(
                    {"hreflang": attrs_map["hreflang"] or "", "href": normalize_url(href, self.page_url) or href}
                )
        elif tag == "a":
            href = attrs_map.get("href")
            if href:
                absolute = urljoin(self.page_url, href)
                self.links.append({"raw": href, "url": normalize_url(absolute, self.page_url) or absolute})
        elif tag == "img":
            self.images.append({"src": attrs_map.get("src"), "alt": attrs_map.get("alt")})
        elif tag == "script" and (attrs_map.get("type") or "").lower() == "application/ld+json":
            self._current_json_ld = []

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "title" and self._text_target and self._text_target[0] == "title":
            self._text_target = None
        elif tag in self.headings and self._current_heading and self._current_heading[0] == tag:
            text = clean_text("".join(self._current_heading[1]))
            if text:
                self.headings[tag].append(text)
            self._current_heading = None
        elif tag == "script" and self._current_json_ld is not None:
            self.json_ld_parts.append(self._current_json_ld)
            self._current_json_ld = None

    def handle_data(self, data: str) -> None:
        if self._text_target:
            self._text_target[1].append(data)
        if self._current_heading:
            self._current_heading[1].append(data)
        if self._current_json_ld is not None:
            self._current_json_ld.append(data)


def inventory_rows() -> list[dict[str, Any]]:
    workbook = load_workbook(INVENTORY, read_only=True, data_only=True)
    sheet = workbook["URL Inventory"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    result = []
    for values in rows[1:]:
        if not values or not values[0] or not str(values[0]).strip().lower().startswith(("http://", "https://")):
            continue
        row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        result.append(row)
    return result


def parse_json_ld(parser: PageParser) -> tuple[list[dict[str, Any]], list[str]]:
    objects: list[dict[str, Any]] = []
    errors: list[str] = []
    for index, parts in enumerate(parser.json_ld_parts, start=1):
        raw = "".join(parts).strip()
        if not raw:
            continue
        try:
            value = json.loads(raw)
            values = value if isinstance(value, list) else [value]
            for item in values:
                if isinstance(item, dict):
                    types = item.get("@type")
                    if isinstance(types, list):
                        type_names = [str(item_type) for item_type in types]
                    elif types:
                        type_names = [str(types)]
                    else:
                        type_names = []
                    objects.append({"script": index, "types": type_names, "id": item.get("@id")})
                else:
                    objects.append({"script": index, "types": [], "id": None})
        except json.JSONDecodeError as error:
            errors.append(f"script {index}: {error.msg} at character {error.pos}")
    return objects, errors


def fetch(url: str, include_text: bool = False) -> dict[str, Any]:
    started = time.monotonic()
    with tempfile.NamedTemporaryFile() as body_file:
        command = [
            "curl", "-sS", "-L", "--max-time", str(TIMEOUT_SECONDS), "--connect-timeout", "10",
            "-A", USER_AGENT, "-H", "Accept: text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "-D", "-", "-o", body_file.name,
            "-w", "\n__STUDIOAETHER_BASELINE__\n%{http_code}\n%{url_effective}\n%{num_redirects}\n%{content_type}\n",
            url,
        ]
        try:
            completed = subprocess.run(command, capture_output=True, timeout=TIMEOUT_SECONDS + 15, check=False)
        except (subprocess.SubprocessError, OSError) as error:
            return {"requested_url": url, "error": str(error), "elapsed_ms": round((time.monotonic() - started) * 1000)}
        body_file.seek(0)
        body = body_file.read()
    marker = b"\n__STUDIOAETHER_BASELINE__\n"
    if marker not in completed.stdout:
        return {"requested_url": url, "error": (completed.stderr or b"curl did not return a status marker").decode("utf-8", errors="replace").strip(), "curl_exit_code": completed.returncode, "elapsed_ms": round((time.monotonic() - started) * 1000)}
    header_bytes, info_bytes = completed.stdout.rsplit(marker, 1)
    info_lines = info_bytes.decode("utf-8", errors="replace").splitlines()
    status = int(info_lines[0]) if info_lines and info_lines[0].isdigit() else None
    final_url = info_lines[1] if len(info_lines) > 1 else url
    content_type = info_lines[3] if len(info_lines) > 3 else ""
    header_blocks = [block for block in re.split(rb"\r?\n\r?\n", header_bytes) if re.search(rb"^HTTP/", block)]
    headers: dict[str, str] = {}
    chain: list[dict[str, Any]] = []
    current_url = url
    for block in header_blocks:
        lines = block.decode("utf-8", errors="replace").splitlines()
        status_match = re.match(r"^HTTP/[^ ]+\s+(\d{3})", lines[0]) if lines else None
        block_status = int(status_match.group(1)) if status_match else None
        block_headers: dict[str, str] = {}
        for line in lines[1:]:
            if ":" in line:
                key, value = line.split(":", 1)
                block_headers[key.lower().strip()] = value.strip()
        if block_status in {301, 302, 303, 307, 308} and block_headers.get("location"):
            next_url = urljoin(current_url, block_headers["location"])
            chain.append({"status": block_status, "from": current_url, "to": next_url})
            current_url = next_url
        if block is header_blocks[-1]:
            headers = block_headers

    text_body = body.decode("utf-8", errors="replace")
    result: dict[str, Any] = {
        "requested_url": url,
        "status": status,
        "final_url": final_url,
        "redirect_chain": chain,
        "redirect_count": len(chain),
        "content_type": content_type,
        "bytes": len(body),
        "sha256": hashlib.sha256(body).hexdigest(),
        "elapsed_ms": round((time.monotonic() - started) * 1000),
        "headers": {name: headers.get(name) for name in ["cache-control", "content-language", "etag", "last-modified", "location", "server", "strict-transport-security", "x-content-type-options", "x-cache-status"] if headers.get(name) is not None},
    }
    if include_text:
        result["text"] = text_body
    if "html" not in content_type.lower() and not text_body.lstrip().startswith("<!doctype html"):
        return result

    parser = PageParser(final_url)
    parser.feed(text_body)
    json_ld, json_ld_errors = parse_json_ld(parser)
    page_url = normalize_url(final_url) or final_url
    internal_origin = urlparse(SITE_ORIGIN).netloc
    internal_links = sorted({link["url"] for link in parser.links if urlparse(link["url"]).netloc == internal_origin})
    external_links = sorted({link["url"] for link in parser.links if urlparse(link["url"]).netloc and urlparse(link["url"]).netloc != internal_origin})
    result["html"] = {
        "lang": parser.lang,
        "title": clean_text("".join(parser.title_parts)) or None,
        "description": clean_text(parser.description or "") or None,
        "robots": clean_text(parser.robots or "") or None,
        "googlebot": clean_text(parser.googlebot or "") or None,
        "canonical": parser.canonical,
        "alternates": parser.alternates,
        "headings": parser.headings,
        "h1_count": len(parser.headings["h1"]),
        "json_ld": json_ld,
        "json_ld_errors": json_ld_errors,
        "json_ld_type_counts": dict(Counter(type_name for item in json_ld for type_name in item["types"])),
        "link_count": len(parser.links),
        "internal_link_count": len(internal_links),
        "external_link_count": len(external_links),
        "internal_links": internal_links,
        "external_links": external_links,
        "image_count": len(parser.images),
        "images_without_alt": sum(1 for image in parser.images if image.get("alt") is None),
        "images_with_empty_alt": sum(1 for image in parser.images if image.get("alt") == ""),
        "normalized_final_url": page_url,
    }
    return result


def fetch_site_file(path: str) -> dict[str, Any]:
    result = fetch(f"{SITE_ORIGIN}{path}", include_text=True)
    body_result = dict(result)
    html = body_result.pop("html", None)
    if html is not None:
        body_result["parsed_as_html"] = True
    return body_result


def comparison(inventory: list[dict[str, Any]], pages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_url = {item["requested_url"]: item for item in pages}
    fields = [("Observed search title", "title"), ("Observed H1 / lead heading", "h1")]
    result = []
    for row in inventory:
        page = by_url.get(row["URL"], {})
        html = page.get("html", {})
        checks: dict[str, Any] = {}
        for inventory_field, html_field in fields:
            expected = clean_text(str(row.get(inventory_field) or ""))
            observed = html.get("title") if html_field == "title" else (html.get("headings", {}).get("h1", [None])[0] if html.get("headings", {}).get("h1") else None)
            checks[inventory_field] = {"inventory": expected or None, "live": observed, "match": expected == (observed or "")}
        result.append({"url": row["URL"], "checks": checks})
    return result


def make_report(report: dict[str, Any]) -> str:
    pages = report["pages"]
    statuses = Counter("error" if page.get("error") else str(page.get("status")) for page in pages)
    redirects = [page for page in pages if page.get("redirect_count", 0)]
    missing = [page for page in pages if page.get("error") or page.get("status") != 200]
    html_pages = [page for page in pages if page.get("html")]
    no_canonical = [page["requested_url"] for page in html_pages if not page["html"].get("canonical")]
    no_description = [page["requested_url"] for page in html_pages if not page["html"].get("description")]
    no_h1 = [page["requested_url"] for page in html_pages if page["html"].get("h1_count") != 1]
    json_errors = [page["requested_url"] for page in html_pages if page["html"].get("json_ld_errors")]
    comparison_rows = report["inventory_comparison"]
    title_mismatches = [row["url"] for row in comparison_rows if not row["checks"]["Observed search title"]["match"]]
    h1_mismatches = [row["url"] for row in comparison_rows if not row["checks"]["Observed H1 / lead heading"]["match"]]
    reciprocal_hreflang = report["checks"]["reciprocal_hreflang"]
    out: list[str] = []
    out.append("# studioaether.com production baseline — 2026-09-04")
    out.append("")
    out.append("Read-only crawl of the current Wix production site, using the committed migration inventory as the URL source. The JSON file contains machine-readable details and payload hashes; full HTML was not committed.")
    out.append("")
    out.append("## Scope")
    out.append("")
    out.append(f"- Crawl timestamp (UTC): `{report['captured_at_utc']}`")
    out.append(f"- Inventory URLs: `{len(report['inventory'])}`")
    out.append(f"- HTML pages parsed: `{len(html_pages)}`")
    out.append(f"- Additional sitemap-discovered URLs crawled: `{len(report.get('discovered_pages', []))}`")
    out.append(f"- Status summary: `{dict(statuses)}`")
    out.append(f"- Site files checked: `/robots.txt`, `/sitemap.xml`, linked English/Hungarian child sitemaps")
    out.append("")
    out.append("## Immediate migration signals")
    out.append("")
    out.append(f"- Non-200 or fetch-error inventory URLs: `{len(missing)}`")
    out.append(f"- Inventory URLs with redirects: `{len(redirects)}`")
    out.append(f"- Pages missing a canonical: `{len(no_canonical)}`")
    out.append(f"- Pages missing a meta description: `{len(no_description)}`")
    out.append(f"- Pages without exactly one H1: `{len(no_h1)}`")
    out.append(f"- Pages with invalid JSON-LD: `{len(json_errors)}`")
    out.append(f"- Live title mismatches against inventory: `{len(title_mismatches)}`")
    out.append(f"- Live first-H1 mismatches against inventory: `{len(h1_mismatches)}`")
    out.append(f"- Reciprocal hreflang issues: `{len(reciprocal_hreflang)}`")
    out.append(f"- Inventory URLs absent from the current XML sitemaps: `{len(report['checks']['sitemap_inventory_urls_missing'])}`")
    out.append(f"- XML sitemap URLs outside the inventory: `{len(report['checks']['sitemap_urls_not_in_inventory'])}`")
    out.append("")
    out.append("## Items requiring deliberate review")
    out.append("")
    out.append("The baseline is evidence for migration decisions, not an instruction to copy every current value. In particular, review the known inventory findings: `/hu/fitness` title, legacy `/mentoring`, indexed `/photographer-daniel`, FAQ pricing/content, public versus legal addresses, and all source-level canonical/hreflang/schema details.")
    out.append("")
    if missing:
        out.append("### Non-200 or errors")
        out.append("")
        for page in missing:
            out.append(f"- `{page['requested_url']}` — status `{page.get('status')}` error `{page.get('error')}` final `{page.get('final_url')}`")
        out.append("")
    if redirects:
        out.append("### Redirecting inventory URLs")
        out.append("")
        for page in redirects:
            chain = " → ".join(item["to"] for item in page["redirect_chain"])
            out.append(f"- `{page['requested_url']}` — `{page['redirect_count']}` hop(s), final `{page['final_url']}`; `{chain}`")
        out.append("")
    if title_mismatches:
        out.append("### Title mismatches")
        out.append("")
        for row in comparison_rows:
            check = row["checks"]["Observed search title"]
            if not check["match"]:
                out.append(f"- `{row['url']}` — inventory `{check['inventory']}`; live `{check['live']}`")
        out.append("")
    if h1_mismatches:
        out.append("### First-H1 mismatches")
        out.append("")
        for row in comparison_rows:
            check = row["checks"]["Observed H1 / lead heading"]
            if not check["match"]:
                out.append(f"- `{row['url']}` — inventory `{check['inventory']}`; live `{check['live']}`")
        out.append("")
    out.append("## Site-level files")
    out.append("")
    for path, details in report["site_files"].items():
        out.append(f"- `{path}` — status `{details.get('status')}`, bytes `{details.get('bytes')}`, SHA-256 `{details.get('sha256')}`")
    out.append("")
    if report.get("discovered_pages"):
        out.append("## Sitemap-discovered URLs outside the inventory")
        out.append("")
        out.append("These URLs were present in the current XML sitemaps but not in the supplied inventory. They must receive an explicit V2 disposition before cutover.")
        out.append("")
        for page in report["discovered_pages"]:
            html = page.get("html", {})
            out.append(f"- `{page['requested_url']}` — status `{page.get('status')}`, final `{page.get('final_url')}`, title `{html.get('title')}`, canonical `{html.get('canonical')}`, H1 count `{html.get('h1_count')}`")
        out.append("")
    out.append("## Per-page details")
    out.append("")
    out.append("See the JSON and CSV files in this directory for the complete per-page record, including metadata, redirect chains, JSON-LD type counts, internal/external link sets, image alt counts, response headers, and payload hashes.")
    out.append("")
    out.append("## Reproduction")
    out.append("")
    out.append("Run from the repository root with the bundled Python runtime:")
    out.append("")
    out.append("```sh")
    out.append("/Users/aczel/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3 scripts/capture-production-baseline.py")
    out.append("```")
    out.append("")
    return "\n".join(out)


def main() -> int:
    inventory = inventory_rows()
    pages: list[dict[str, Any]] = []
    for index, row in enumerate(inventory):
        if index:
            time.sleep(REQUEST_DELAY_SECONDS)
        print(f"[{index + 1}/{len(inventory)}] {row['URL']}", file=sys.stderr)
        pages.append(fetch(row["URL"]))

    site_files = {}
    for path in ("/robots.txt", "/sitemap.xml", "/pages-sitemap.xml", "/hu_hu-sitemap.xml", "/hu_hu-pages-sitemap.xml"):
        time.sleep(REQUEST_DELAY_SECONDS)
        site_files[path] = fetch_site_file(path)

    for path, details in site_files.items():
        text_body = details.get("text", "")
        if path == "/robots.txt":
            details["parsed"] = {
                "sitemap_directives": re.findall(r"(?im)^\s*sitemap:\s*(\S+)", text_body),
                "disallow_directives": re.findall(r"(?im)^\s*disallow:\s*(\S*)", text_body),
                "has_global_allow": bool(re.search(r"(?im)^\s*allow:\s*/\s*$", text_body)),
            }
        elif path.endswith(".xml"):
            details["parsed"] = {
                "locs": re.findall(r"<loc>\s*(.*?)\s*</loc>", text_body, flags=re.I | re.S),
                "lastmods": re.findall(r"<lastmod>\s*(.*?)\s*</lastmod>", text_body, flags=re.I | re.S),
                "is_sitemap_index": "<sitemapindex" in text_body.lower(),
                "is_urlset": "<urlset" in text_body.lower(),
            }
        details.pop("text", None)

    inventory_urls = {normalize_url(row["URL"]) for row in inventory}
    sitemap_urls = set()
    for details in site_files.values():
        parsed = details.get("parsed", {})
        if not parsed.get("is_urlset"):
            continue
        for loc in parsed.get("locs", []):
            normalized = normalize_url(loc)
            if normalized:
                sitemap_urls.add(normalized)

    extra_urls = sorted(sitemap_urls - inventory_urls)
    discovered_pages: list[dict[str, Any]] = []
    for url in extra_urls:
        time.sleep(REQUEST_DELAY_SECONDS)
        print(f"[sitemap extra {len(discovered_pages) + 1}/{len(extra_urls)}] {url}", file=sys.stderr)
        discovered_pages.append(fetch(url))

    by_url = {page["requested_url"]: page for page in pages}
    expected_pairs = {
        row["URL"]: row.get("Counterpart")
        for row in inventory
        if str(row.get("Counterpart") or "").strip().lower().startswith(("http://", "https://"))
    }
    hreflang_issues = []
    normalized_page_by_url = {normalize_url(page["requested_url"]): page for page in pages}
    for page in pages:
        html = page.get("html", {})
        alternates = {item["href"] for item in html.get("alternates", [])}
        current_url = normalize_url(page.get("final_url") or page["requested_url"])
        counterpart = normalize_url(expected_pairs.get(page["requested_url"]))
        if counterpart and counterpart not in alternates:
            hreflang_issues.append({"url": page["requested_url"], "direction": "outbound", "expected_counterpart": counterpart, "alternates": html.get("alternates", [])})
        counterpart_page = normalized_page_by_url.get(counterpart) if counterpart else None
        if counterpart_page and current_url:
            counterpart_alternates = {item["href"] for item in counterpart_page.get("html", {}).get("alternates", [])}
            if current_url not in counterpart_alternates:
                hreflang_issues.append({"url": page["requested_url"], "direction": "inbound", "counterpart": counterpart, "counterpart_alternates": counterpart_page.get("html", {}).get("alternates", [])})

    internal_link_issues = []
    for page in pages:
        for target in page.get("html", {}).get("internal_links", []):
            if target in inventory_urls:
                target_page = by_url.get(next((row["URL"] for row in inventory if normalize_url(row["URL"]) == target), ""), {})
                if target_page.get("status") != 200 or target_page.get("redirect_count", 0):
                    internal_link_issues.append({"source": page["requested_url"], "target": target, "target_status": target_page.get("status"), "target_redirect_count": target_page.get("redirect_count")})

    report = {
        "captured_at_utc": datetime.now(timezone.utc).isoformat(),
        "site_origin": SITE_ORIGIN,
        "user_agent": USER_AGENT,
        "inventory_source": str(INVENTORY.relative_to(ROOT)),
        "inventory": inventory,
        "pages": pages,
        "discovered_pages": discovered_pages,
        "site_files": site_files,
        "inventory_comparison": comparison(inventory, pages),
        "checks": {
            "reciprocal_hreflang": hreflang_issues,
            "internal_links_to_non200_or_redirecting_inventory_urls": internal_link_issues,
            "sitemap_inventory_urls_missing": sorted(inventory_urls - sitemap_urls),
            "sitemap_urls_not_in_inventory": sorted(sitemap_urls - inventory_urls),
        },
    }
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    json_path = OUTPUT_DIR / "production-baseline-2026-09-04.json"
    markdown_path = OUTPUT_DIR / "production-baseline-2026-09-04.md"
    csv_path = OUTPUT_DIR / "production-page-baseline-2026-09-04.csv"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    markdown_path.write_text(make_report(report), encoding="utf-8")
    with csv_path.open("w", newline="", encoding="utf-8") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=[
            "url", "status", "final_url", "redirect_count", "title", "description", "lang",
            "canonical", "hreflang", "h1_count", "h1s", "h2_count", "robots", "googlebot",
            "json_ld_types", "json_ld_errors", "internal_link_count", "external_link_count",
            "image_count", "images_without_alt", "images_with_empty_alt", "bytes", "sha256",
        ])
        writer.writeheader()
        for page in pages:
            html = page.get("html", {})
            writer.writerow({
                "url": page.get("requested_url"),
                "status": page.get("status"),
                "final_url": page.get("final_url"),
                "redirect_count": page.get("redirect_count"),
                "title": html.get("title"),
                "description": html.get("description"),
                "lang": html.get("lang"),
                "canonical": html.get("canonical"),
                "hreflang": "; ".join(f"{item.get('hreflang')}={item.get('href')}" for item in html.get("alternates", [])),
                "h1_count": html.get("h1_count"),
                "h1s": " | ".join(html.get("headings", {}).get("h1", [])),
                "h2_count": len(html.get("headings", {}).get("h2", [])),
                "robots": html.get("robots"),
                "googlebot": html.get("googlebot"),
                "json_ld_types": "; ".join(f"{key}={value}" for key, value in sorted(html.get("json_ld_type_counts", {}).items())),
                "json_ld_errors": " | ".join(html.get("json_ld_errors", [])),
                "internal_link_count": html.get("internal_link_count"),
                "external_link_count": html.get("external_link_count"),
                "image_count": html.get("image_count"),
                "images_without_alt": html.get("images_without_alt"),
                "images_with_empty_alt": html.get("images_with_empty_alt"),
                "bytes": page.get("bytes"),
                "sha256": page.get("sha256"),
            })
    print(f"Wrote {json_path}", file=sys.stderr)
    print(f"Wrote {markdown_path}", file=sys.stderr)
    print(f"Wrote {csv_path}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
